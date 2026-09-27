from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_ppt():
    prs = Presentation()
    
    # Title Slide
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Soil Health Prediction\nUsing Machine Learning"
    subtitle.text = "A Full-Stack Data Science & Engineering Project\n\nAuthor: [Your Name/Team Name]"
    
    # Introduction
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Introduction & Problem Statement"
    content = slide.placeholders[1]
    content.text = ("• The Agricultural Challenge:\n"
                    "  - Traditional soil testing is manual, slow, and expensive.\n"
                    "  - Delayed results lead to suboptimal fertilizer use and reduced crop yields.\n\n"
                    "• The Solution:\n"
                    "  - A real-time, machine learning-driven software application.\n"
                    "  - Instant analysis of soil parameters (pH, Moisture, N-P-K) to predict overall soil health.")
    
    # Objectives
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Project Objectives"
    content = slide.placeholders[1]
    content.text = ("1. Develop a high-accuracy ML classification model to assess soil quality.\n"
                    "2. Architect a scalable Java Spring Boot backend to serve the model.\n"
                    "3. Design an intuitive, responsive web dashboard for end-users.\n"
                    "4. Create a seamless end-to-end data pipeline from user input to ML inference and back.")
    
    # Architecture
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "System Architecture"
    content = slide.placeholders[1]
    content.text = ("• Frontend (Client Layer):\n"
                    "  - Modern Web Dashboard (HTML5, CSS3, JS).\n"
                    "  - Captures sensor/manual data.\n\n"
                    "• Backend (Application Layer):\n"
                    "  - Java Spring Boot REST API.\n"
                    "  - MongoDB Stitch for database connectivity and logging.\n\n"
                    "• Machine Learning (Inference Layer):\n"
                    "  - ONNX Runtime embedded in Java for ultra-fast, native predictions.")
    
    # Tech Stack
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Technology Stack"
    content = slide.placeholders[1]
    content.text = ("• Data Science & Modeling:\n"
                    "  - Python, Pandas, Numpy\n"
                    "  - Scikit-Learn (Random Forest)\n"
                    "  - ONNX (Open Neural Network Exchange)\n\n"
                    "• Software Engineering:\n"
                    "  - Java 17, Spring Boot\n"
                    "  - MongoDB Stitch\n"
                    "  - Frontend Web Technologies (HTML, CSS, JS, Framer Motion principles)")
    
    # Data Collection & Preprocessing
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Dataset & Preprocessing"
    content = slide.placeholders[1]
    content.text = ("• Dataset Overview:\n"
                    "  - 1,000 comprehensive soil samples.\n"
                    "  - Features: pH (4-9), Moisture (10-90%), Nitrogen, Phosphorus, Potassium.\n\n"
                    "• Preprocessing Steps:\n"
                    "  - Outlier detection and data normalization.\n"
                    "  - Stratified Train-Test Split (80% Training, 20% Testing).\n"
                    "  - Label encoding for soil quality: Poor, Average, Good.")
    
    # Machine Learning Methodology
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Machine Learning Methodology"
    content = slide.placeholders[1]
    content.text = ("• Algorithm Selected: Random Forest Classifier\n"
                    "  - Why? High accuracy, robustness against overfitting, and capability to handle non-linear environmental data.\n\n"
                    "• Training & Export:\n"
                    "  - The model was trained using Scikit-Learn.\n"
                    "  - Serialized to the .onnx format to completely decouple the Python training environment from the Java production environment.")
    
    # User Interface
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "User Interface Design"
    content = slide.placeholders[1]
    content.text = ("• Design Philosophy:\n"
                    "  - Clean, modern, and data-centric.\n"
                    "  - Responsive layout ensuring accessibility on both desktop and mobile devices.\n\n"
                    "• Key Features:\n"
                    "  - Real-time form validation.\n"
                    "  - Asynchronous loading states and dynamic visual feedback (color-coded health rings).")
    
    # Results
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Results & Performance"
    content = slide.placeholders[1]
    content.text = ("• Model Performance:\n"
                    "  - Achieved ~96% classification accuracy on the testing set.\n"
                    "  - High precision and recall across all three soil health categories.\n\n"
                    "• System Performance:\n"
                    "  - Spring Boot API inference latency < 50ms.\n"
                    "  - Highly stable integration between Java and the ONNX ML model.")
    
    # Conclusion
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Conclusion"
    content = slide.placeholders[1]
    content.text = ("• The project successfully bridges Data Science and Software Engineering.\n"
                    "• Proves that real-time, AI-driven agricultural insights can be delivered through a scalable web platform.\n"
                    "• Empowers agricultural stakeholders to make rapid, data-driven decisions regarding soil management.")
    
    # Future Scope
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Future Scope"
    content = slide.placeholders[1]
    content.text = ("1. Hardware Integration:\n"
                    "   - Connecting physical IoT soil sensors for automated data ingestion.\n\n"
                    "2. Advanced Modeling:\n"
                    "   - Incorporating weather and historical yield data to recommend specific crop types.\n\n"
                    "3. Mobile Deployment:\n"
                    "   - Developing a native mobile application (e.g., using React Native) for field workers.")
                    
    # Q&A
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Thank You"
    subtitle.text = "Questions & Discussion"
    
    prs.save('Project_Presentation.pptx')
    print("Professional Presentation generated successfully.")

if __name__ == '__main__':
    create_ppt()
